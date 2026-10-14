"""
Router Artículo 22 — Control de Presencia en Planta
Empleados excluidos de control de asistencia (Art. 22 Código del Trabajo Chile)
"""
from fastapi import APIRouter, Depends, HTTPException, Query
from typing import List, Dict, Any, Optional
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo
from loguru import logger

from backend.core.config import settings

def _now_chile():
    """Retorna datetime naive en hora de Chile (respeta DST automáticamente)"""
    tz = ZoneInfo(settings.TIMEZONE)  # "America/Santiago"
    return datetime.now(tz).replace(tzinfo=None)

from backend.core.database import get_db, Database
from backend.core.security import SecurityContext, get_current_user, RequirePermission

router = APIRouter(
    prefix="/articulo22",
    tags=["Artículo 22"]
)


# ============================================
# INICIALIZACIÓN DE TABLA
# ============================================
async def _ensure_table(db: Database):
    """Crea la tabla si no existe."""
    await db.execute("""
        CREATE TABLE IF NOT EXISTS articulo22_registros (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empleado_id INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            hora TEXT NOT NULL,
            tipo TEXT NOT NULL,
            registrado_por_id INTEGER,
            registrado_por_nombre TEXT,
            observaciones TEXT DEFAULT '',
            created_at TEXT DEFAULT (datetime('now')),
            FOREIGN KEY (empleado_id) REFERENCES empleados(id)
        )
    """)
    # Índices (CREATE INDEX IF NOT EXISTS no falla si ya existen)
    await db.execute("CREATE INDEX IF NOT EXISTS idx_art22_fecha ON articulo22_registros(fecha)")
    await db.execute("CREATE INDEX IF NOT EXISTS idx_art22_emp_fecha ON articulo22_registros(empleado_id, fecha)")


def _estadia_min(marcas: List[Dict[str, Any]]) -> float:
    """Minutos en planta de un día: cada ENTRADA seguida de una SALIDA suma su diferencia (cruza medianoche si S < E)."""
    total = 0.0
    i = 0
    while i < len(marcas):
        m = marcas[i]
        if m["tipo"] == "E" and i + 1 < len(marcas) and marcas[i + 1]["tipo"] == "S":
            e_dt = datetime.strptime(m["hora"], "%H:%M:%S")
            s_dt = datetime.strptime(marcas[i + 1]["hora"], "%H:%M:%S")
            if s_dt < e_dt:
                s_dt += timedelta(days=1)
            total += (s_dt - e_dt).total_seconds() / 60
            i += 2
            continue
        i += 1
    return total


def _fmt_hm(minutos: float) -> str:
    m = int(round(minutos))
    return f"{m // 60}h {m % 60:02d}m"


# ============================================
# ENDPOINTS
# ============================================

@router.get("/empleados/")
async def get_empleados_art22(
    db: Database = Depends(get_db),
    current_user: SecurityContext = Depends(get_current_user),
    _perm = Depends(RequirePermission("porteria.art22"))
):
    """Lista empleados con excluido_asistencia=1 incluyendo datos de ficha."""
    rows = await db.fetch_all("""
        SELECT e.id, e.rut, e.nombre, e.apellido_paterno, e.apellido_materno,
               e.cargo, a.nombre as area_nombre, e.activo
        FROM empleados e
        LEFT JOIN areas a ON e.area_id = a.id
        WHERE e.excluido_asistencia = 1 AND e.activo = 1
        ORDER BY e.apellido_paterno, e.nombre
    """)
    return [dict(r) for r in rows]


@router.post("/marcar/")
async def marcar_presencia(
    data: Dict[str, Any],
    db: Database = Depends(get_db),
    current_user: SecurityContext = Depends(get_current_user),
    _perm = Depends(RequirePermission("porteria.art22"))
):
    """
    Registra una marca de presencia. Auto-detecta tipo E/S.
    Body: { "empleado_id": int, "observaciones": str (opcional) }
    """
    await _ensure_table(db)

    empleado_id = data.get("empleado_id")
    observaciones = data.get("observaciones", "")

    if not empleado_id:
        raise HTTPException(status_code=400, detail="empleado_id requerido")

    # Verificar que sea empleado Art. 22
    emp = await db.fetch_one(
        "SELECT id, nombre, apellido_paterno, excluido_asistencia FROM empleados WHERE id = ?",
        (empleado_id,)
    )
    if not emp:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")
    if not emp.get("excluido_asistencia"):
        raise HTTPException(status_code=400, detail="Este empleado no tiene Artículo 22")

    # Fecha y hora actual (Chile)
    now = _now_chile()
    fecha_hoy = now.strftime("%Y-%m-%d")
    hora_actual = now.strftime("%H:%M:%S")

    # Obtener última marca GLOBAL del empleado (sin filtro de fecha)
    # para detectar correctamente E/S en sesiones cross-midnight
    ultima = await db.fetch_one(
        "SELECT tipo, fecha, hora FROM articulo22_registros WHERE empleado_id = ? ORDER BY id DESC LIMIT 1",
        (empleado_id,)
    )

    # Protección contra doble click / debounce en el backend (ej. 5 segundos)
    if ultima:
        try:
            last_dt = datetime.strptime(f"{ultima['fecha']} {ultima['hora']}", "%Y-%m-%d %H:%M:%S")
            time_diff = (now - last_dt).total_seconds()
            if time_diff < 5:
                logger.warning(f"[ART22] Ignorando marca duplicada para empleado {empleado_id} (diferencia: {time_diff}s)")
                nombre_emp = f"{emp.get('apellido_paterno', '')} {emp.get('nombre', '')}"
                tipo_label = "ENTRADA" if ultima["tipo"] == "E" else "SALIDA"
                return {
                    "ok": True,
                    "tipo": ultima["tipo"],
                    "tipo_label": tipo_label,
                    "hora": ultima["hora"],
                    "empleado_id": empleado_id,
                    "nombre": nombre_emp
                }
        except Exception as ex:
            logger.error(f"[ART22] Error comparando tiempos de duplicación: {ex}")

    nuevo_tipo = "E" if (not ultima or ultima["tipo"] == "S") else "S"

    # Determinar fecha de registro:
    # Si es Salida y la Entrada fue de otro día → guardar con fecha de la Entrada
    # para mantener la sesión completa en el mismo día
    if nuevo_tipo == "S" and ultima and ultima["fecha"] != fecha_hoy:
        fecha_registro = ultima["fecha"]  # Cross-midnight: atribuir al día de entrada
        logger.info(f"[ART22] Cross-midnight detectado: Salida de hoy atribuida a {fecha_registro}")
    else:
        fecha_registro = fecha_hoy

    # Registrar
    registrador = f"{current_user.nombre_completo}" if hasattr(current_user, 'nombre_completo') else (current_user.username or "Sistema")

    await db.execute(
        """INSERT INTO articulo22_registros 
           (empleado_id, fecha, hora, tipo, registrado_por_id, registrado_por_nombre, observaciones)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        (empleado_id, fecha_registro, hora_actual, nuevo_tipo, 
         current_user.user_id if hasattr(current_user, 'user_id') else None,
         registrador, observaciones)
    )

    nombre_emp = f"{emp.get('apellido_paterno', '')} {emp.get('nombre', '')}"
    tipo_label = "ENTRADA" if nuevo_tipo == "E" else "SALIDA"
    logger.info(f"[ART22] {tipo_label}: {nombre_emp} (ID={empleado_id}) a las {hora_actual} - Reg: {registrador}")

    return {
        "ok": True,
        "tipo": nuevo_tipo,
        "tipo_label": tipo_label,
        "hora": hora_actual,
        "empleado_id": empleado_id,
        "nombre": nombre_emp
    }


@router.get("/estado-dia/")
async def get_estado_dia(
    fecha: Optional[str] = Query(None, description="Fecha YYYY-MM-DD, default hoy"),
    db: Database = Depends(get_db),
    current_user: SecurityContext = Depends(get_current_user),
    _perm = Depends(RequirePermission("porteria.art22"))
):
    """
    Estado de presencia de todos los empleados Art. 22 para un día.
    Incluye marcas, estadía acumulada, y estado actual.
    """
    await _ensure_table(db)

    if not fecha:
        fecha = _now_chile().strftime("%Y-%m-%d")

    # Obtener empleados Art. 22
    empleados = await db.fetch_all("""
        SELECT e.id, e.rut, e.nombre, e.apellido_paterno, e.apellido_materno,
               e.cargo, a.nombre as area_nombre
        FROM empleados e
        LEFT JOIN areas a ON e.area_id = a.id
        WHERE e.excluido_asistencia = 1 AND e.activo = 1
        ORDER BY e.apellido_paterno, e.nombre
    """)

    # Obtener todas las marcas del día
    marcas = await db.fetch_all(
        """SELECT empleado_id, hora, tipo, registrado_por_nombre, observaciones
           FROM articulo22_registros
           WHERE fecha = ?
           ORDER BY empleado_id, id""",
        (fecha,)
    )

    # Agrupar marcas por empleado
    marcas_por_emp = {}
    for m in marcas:
        eid = m["empleado_id"]
        if eid not in marcas_por_emp:
            marcas_por_emp[eid] = []
        marcas_por_emp[eid].append(dict(m))

    # Calcular estado por empleado
    resultado = []
    now = _now_chile()
    es_hoy = (fecha == now.strftime("%Y-%m-%d"))

    for emp in empleados:
        eid = emp["id"]
        emp_marcas = marcas_por_emp.get(eid, [])

        nombre_completo = f"{emp.get('apellido_paterno', '')} {emp.get('apellido_materno', '')}, {emp.get('nombre', '')}"

        if not emp_marcas:
            resultado.append({
                "empleado_id": eid,
                "nombre": nombre_completo.strip(", "),
                "cargo": emp.get("cargo", ""),
                "area": emp.get("area_nombre", ""),
                "marcas": [],
                "primera_entrada": None,
                "ultima_marca": None,
                "sesion_abierta": False,
                "estado": "sin_registro",
                "estadia_total_min": 0,
                "estadia_display": "—"
            })
            continue

        # Calcular estadía acumulada
        primera_entrada = None
        ultima_marca = emp_marcas[-1]
        sesion_abierta = (ultima_marca["tipo"] == "E")
        estadia_min = 0

        i = 0
        while i < len(emp_marcas):
            m = emp_marcas[i]
            if m["tipo"] == "E":
                if primera_entrada is None:
                    primera_entrada = m["hora"]
                # Buscar la salida correspondiente
                entrada_dt = datetime.strptime(f"{fecha} {m['hora']}", "%Y-%m-%d %H:%M:%S")
                if i + 1 < len(emp_marcas) and emp_marcas[i + 1]["tipo"] == "S":
                    salida_dt = datetime.strptime(f"{fecha} {emp_marcas[i + 1]['hora']}", "%Y-%m-%d %H:%M:%S")
                    if salida_dt < entrada_dt:
                        salida_dt += timedelta(days=1)  # Cross-midnight
                    estadia_min += (salida_dt - entrada_dt).total_seconds() / 60
                    i += 2
                    continue
                elif es_hoy and sesion_abierta and i == len(emp_marcas) - 1:
                    # Sesión abierta, calcular hasta ahora
                    estadia_min += (now - entrada_dt).total_seconds() / 60
            i += 1

        horas = int(estadia_min // 60)
        mins = int(estadia_min % 60)
        display = f"{horas}h {mins:02d}m"
        if sesion_abierta and es_hoy:
            display += " ⏳"

        estado = "sin_registro"
        if emp_marcas:
            estado = "en_planta" if sesion_abierta else "fuera"

        resultado.append({
            "empleado_id": eid,
            "nombre": nombre_completo.strip(", "),
            "cargo": emp.get("cargo", ""),
            "area": emp.get("area_nombre", ""),
            "marcas": emp_marcas,
            "primera_entrada": primera_entrada,
            "ultima_marca": ultima_marca["hora"] if ultima_marca else None,
            "ultima_marca_tipo": ultima_marca["tipo"] if ultima_marca else None,
            "sesion_abierta": sesion_abierta,
            "estado": estado,
            "estadia_total_min": round(estadia_min, 1),
            "estadia_display": display
        })

    # Stats
    total = len(resultado)
    con_registro = sum(1 for r in resultado if r["estado"] != "sin_registro")
    en_planta = sum(1 for r in resultado if r["estado"] == "en_planta")

    return {
        "fecha": fecha,
        "stats": {
            "total": total,
            "con_registro": con_registro,
            "en_planta": en_planta,
            "sin_registro": total - con_registro
        },
        "empleados": resultado
    }


@router.get("/historial/")
async def get_historial(
    desde: Optional[str] = Query(None),
    hasta: Optional[str] = Query(None),
    area: Optional[str] = Query(None),
    empleado_id: Optional[int] = Query(None),
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),
    db: Database = Depends(get_db),
    current_user: SecurityContext = Depends(get_current_user),
    _perm = Depends(RequirePermission("porteria.art22"))
):
    """Historial de registros Art. 22 con filtros y paginación."""
    await _ensure_table(db)

    if not hasta:
        hasta = _now_chile().strftime("%Y-%m-%d")
    if not desde:
        desde = (datetime.strptime(hasta, "%Y-%m-%d") - timedelta(days=7)).strftime("%Y-%m-%d")

    # Consulta agrupada por empleado+fecha
    where_clauses = ["r.fecha BETWEEN ? AND ?"]
    params = [desde, hasta]

    if area:
        where_clauses.append("a.nombre = ?")
        params.append(area)
    if empleado_id:
        where_clauses.append("r.empleado_id = ?")
        params.append(empleado_id)

    where_sql = " AND ".join(where_clauses)

    # Obtener resumen por empleado/día
    count_row = await db.fetch_one(
        f"""SELECT COUNT(DISTINCT r.empleado_id || r.fecha) as total
            FROM articulo22_registros r
            JOIN empleados e ON r.empleado_id = e.id
            LEFT JOIN areas a ON e.area_id = a.id
            WHERE {where_sql}""",
        tuple(params)
    )
    total_records = count_row["total"] if count_row else 0

    offset = (page - 1) * limit
    params_query = list(params) + [limit, offset]

    rows = await db.fetch_all(
        f"""SELECT r.empleado_id, r.fecha,
                   e.nombre as emp_nombre, e.apellido_paterno, e.apellido_materno,
                   e.cargo, a.nombre as area_nombre,
                   MIN(CASE WHEN r.tipo='E' THEN r.hora END) as primera_entrada,
                   MAX(CASE WHEN r.tipo='S' THEN r.hora END) as ultima_salida,
                   COUNT(*) as total_marcas
            FROM articulo22_registros r
            JOIN empleados e ON r.empleado_id = e.id
            LEFT JOIN areas a ON e.area_id = a.id
            WHERE {where_sql}
            GROUP BY r.empleado_id, r.fecha
            ORDER BY r.fecha DESC, e.apellido_paterno
            LIMIT ? OFFSET ?""",
        tuple(params_query)
    )

    # Para cada fila, calcular estadía
    resultado = []
    for row in rows:
        r = dict(row)
        # Obtener marcas detalladas
        marcas = await db.fetch_all(
            "SELECT hora, tipo FROM articulo22_registros WHERE empleado_id = ? AND fecha = ? ORDER BY id",
            (r["empleado_id"], r["fecha"])
        )
        estadia_min = _estadia_min([dict(m) for m in marcas])

        horas = int(estadia_min // 60)
        mins = int(estadia_min % 60)
        nombre = f"{r.get('apellido_paterno', '')} {r.get('apellido_materno', '')}, {r.get('emp_nombre', '')}"

        resultado.append({
            "fecha": r["fecha"],
            "empleado_id": r["empleado_id"],
            "nombre": nombre.strip(", "),
            "cargo": r.get("cargo", ""),
            "area": r.get("area_nombre", ""),
            "primera_entrada": r.get("primera_entrada"),
            "ultima_salida": r.get("ultima_salida"),
            "total_marcas": r.get("total_marcas", 0),
            "estadia_min": round(estadia_min, 1),
            "estadia_display": f"{horas}h {mins:02d}m"
        })

    # Resumen de TODO el rango para un solo empleado (no solo la página visible)
    resumen = None
    if empleado_id:
        todas = await db.fetch_all(
            "SELECT fecha, hora, tipo FROM articulo22_registros WHERE empleado_id = ? AND fecha BETWEEN ? AND ? ORDER BY fecha, id",
            (empleado_id, desde, hasta)
        )
        por_dia: Dict[str, List[Dict[str, Any]]] = {}
        for m in todas:
            por_dia.setdefault(m["fecha"], []).append(dict(m))
        total_min = sum(_estadia_min(v) for v in por_dia.values())
        dias_con = len(por_dia)
        dias_incompletos = sum(
            1 for v in por_dia.values()
            if sum(1 for x in v if x["tipo"] == "E") != sum(1 for x in v if x["tipo"] == "S")
        )
        resumen = {
            "dias": dias_con,
            "estadia_min": round(total_min, 1),
            "estadia_display": _fmt_hm(total_min),
            "horas_decimal": round(total_min / 60, 2),
            "promedio_dia_display": _fmt_hm(total_min / dias_con) if dias_con else "0h 00m",
            "primera_fecha": min(por_dia) if por_dia else None,
            "ultima_fecha": max(por_dia) if por_dia else None,
            "dias_incompletos": dias_incompletos,
        }

    return {
        "resumen": resumen,
        "desde": desde,
        "hasta": hasta,
        "total": total_records,
        "page": page,
        "limit": limit,
        "pages": (total_records + limit - 1) // limit,
        "registros": resultado
    }


@router.get("/exportar-excel/")
async def exportar_excel(
    desde: str = Query(...),
    hasta: str = Query(...),
    empleado_id: Optional[int] = Query(None),
    area: Optional[str] = Query(None),
    db: Database = Depends(get_db),
    current_user: SecurityContext = Depends(get_current_user),
    _perm = Depends(RequirePermission("porteria.art22"))
):
    """Excel con las marcaciones Art. 22 del rango (todos o un empleado): resumen por día, marcas y total por empleado."""
    import io
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    await _ensure_table(db)
    try:
        d0 = datetime.strptime(desde, "%Y-%m-%d")
        d1 = datetime.strptime(hasta, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail="Fechas inválidas (use AAAA-MM-DD)")
    if d1 < d0:
        raise HTTPException(status_code=400, detail="La fecha final es anterior a la inicial")
    if (d1 - d0).days > 3660:
        raise HTTPException(status_code=400, detail="El rango máximo es de 10 años")

    where = ["r.fecha BETWEEN ? AND ?"]
    params: List[Any] = [desde, hasta]
    if empleado_id:
        where.append("r.empleado_id = ?")
        params.append(empleado_id)
    if area:
        where.append("a.nombre = ?")
        params.append(area)
    rows = await db.fetch_all(
        f"""SELECT r.empleado_id, r.fecha, r.hora, r.tipo, r.registrado_por_nombre, r.observaciones,
                   e.rut, e.nombre, e.apellido_paterno, e.apellido_materno, e.cargo, a.nombre AS area_nombre
            FROM articulo22_registros r
            JOIN empleados e ON r.empleado_id = e.id
            LEFT JOIN areas a ON e.area_id = a.id
            WHERE {" AND ".join(where)}
            ORDER BY e.apellido_paterno, e.nombre, r.fecha, r.id""",
        tuple(params)
    )
    if not rows:
        raise HTTPException(status_code=404, detail="No hay marcaciones en el rango seleccionado")

    emps: Dict[int, Dict[str, Any]] = {}
    for r in rows:
        r = dict(r)
        e = emps.setdefault(r["empleado_id"], {
            "rut": r["rut"],
            "nombre": f"{r['apellido_paterno'] or ''} {r['apellido_materno'] or ''}, {r['nombre'] or ''}".strip(" ,"),
            "cargo": r["cargo"] or "", "area": r["area_nombre"] or "", "dias": {}})
        e["dias"].setdefault(r["fecha"], []).append(r)

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1E3A8A")

    def _encabezado(ws, cols, anchos):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center", vertical="center")
        for i, w in enumerate(anchos, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        ws.freeze_panes = "A2"

    ws1 = wb.active
    ws1.title = "Resumen por día"
    _encabezado(ws1, ["Fecha", "RUT", "Empleado", "Cargo", "Área", "1ª entrada", "Últ. salida", "Marcas",
                      "Estadía (h:mm)", "Horas (decimal)", "Observación"],
                [12, 13, 34, 30, 22, 11, 11, 8, 14, 14, 36])
    ws2 = wb.create_sheet("Marcaciones")
    _encabezado(ws2, ["Fecha", "RUT", "Empleado", "Hora", "Tipo", "Registrado por", "Observaciones"],
                [12, 13, 34, 10, 10, 28, 40])
    ws3 = wb.create_sheet("Total por empleado")
    _encabezado(ws3, ["RUT", "Empleado", "Cargo", "Área", "Primera fecha", "Última fecha", "Días con marcas",
                      "Estadía total (h:mm)", "Horas (decimal)", "Promedio por día", "Días con marca sin par"],
                [13, 34, 30, 22, 13, 13, 15, 18, 15, 16, 20])

    gran_total = 0.0
    for e in emps.values():
        tot = 0.0
        incompletos = 0
        for fecha in sorted(e["dias"]):
            ms = [{"hora": m["hora"], "tipo": m["tipo"]} for m in e["dias"][fecha]]
            mins = _estadia_min(ms)
            tot += mins
            ents = [m["hora"] for m in ms if m["tipo"] == "E"]
            sals = [m["hora"] for m in ms if m["tipo"] == "S"]
            aviso = ""
            if len(ents) != len(sals):
                incompletos += 1
                aviso = "Entrada o salida sin par: no suma horas"
            ws1.append([fecha, e["rut"], e["nombre"], e["cargo"], e["area"], ents[0][:5] if ents else "",
                        sals[-1][:5] if sals else "", len(ms), _fmt_hm(mins), round(mins / 60, 2), aviso])
            for m in e["dias"][fecha]:
                ws2.append([fecha, e["rut"], e["nombre"], m["hora"], "Entrada" if m["tipo"] == "E" else "Salida",
                            m["registrado_por_nombre"] or "", m["observaciones"] or ""])
        gran_total += tot
        dias = len(e["dias"])
        ws3.append([e["rut"], e["nombre"], e["cargo"], e["area"], min(e["dias"]), max(e["dias"]), dias,
                    _fmt_hm(tot), round(tot / 60, 2), _fmt_hm(tot / dias) if dias else "", incompletos])
    if len(emps) > 1:
        ws3.append([])
        ws3.append(["", "TOTAL", "", "", "", "", "", _fmt_hm(gran_total), round(gran_total / 60, 2)])
        for c in ws3[ws3.max_row]:
            c.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    sufijo = f"_emp{empleado_id}" if empleado_id else ""
    nombre = f"articulo22_{desde}_a_{hasta}{sufijo}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'}
    )
